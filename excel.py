# '''
# **openpyxl** es la librería estándar de Pyhton para leer, crear y modificar archivos de Excel en formato moderno (`.xlsx). Permite
# desde tareas sencillas como escribir datos en celdas hasta automatizaciones avanzadas como aplicar estilos, formatos condicionales, 
# fórmulas y generar gráficos. 
# Conceptos claves:

# Workbook (El libro completo de Excel).
# Worksheet (Una hoja específica del libro).
# Cell (La celda individual donde guaradas valores o fórmulas).
# '''

# from openpyxl import Workbook

# wb = Workbook()

# hoja_1 = wb.active # Tiene indice 0
# hoja_1.title = "Hoja 1"

# hoja_2 = wb.create_sheet('Mi tercera hoja')

# hoja_3 = wb.create_sheet('Mi segunda hoja', 1)

# print(wb.sheetnames)

# hoja_1['A1'] = 'Productos'
# hoja_1['B1'] = 'Precio'

# hoja_1.cell(row=2, column=1, value='Laptop')
# hoja_1.cell(row=2, column=2, value=1500)

# wb.save('Ventas.xlsx')

from openpyxl import Workbook

wb = Workbook()
ws = wb.active
ws.title = "Inventario"

ws.cell(row=1, column=1, value='Producto')
ws.cell(row=1, column=2, value='Precio')

productos = [ ("Laptop Pro 15", 1200), ("Mouse Inalámbrico", 25), ("Teclado Mecánico", 45), ("Monitor 4K 27'", 350), ("Impresora Láser", 180), ("Auriculares Bluetooth", 85), ("Disco Duro Externo 1TB", 60), ("Cámara Web HD", 50), ("Silla Ergonómica", 220), ("Escritorio Elevable", 400) ]

for fila, (prod, precio) in enumerate(productos, start=2): 
    ws.cell(row=fila, column=1, value=prod) # Columna 1 (A) 
    ws.cell(row=fila, column=2, value=precio) # Columna 2 (B)
    print(f"Iteración: fila={fila}, Producto={prod}, Precio={precio}")

wb.save('Inventario_tienda.xlsx')

for fila in range(2,11):
    prod = ws.cell(row=fila, column=1).value
    precio = ws.cell(row=fila, column=2).value
    print(f"Fila {fila}: Producto={prod}, Precio={precio}")